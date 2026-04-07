from __future__ import annotations
import io
import zipfile
from typing import TYPE_CHECKING

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app.pipeline.renderers.manual_renderer import (
    IRElement, IRHeading, IRParagraph, IRList, IRTable,
)
from app.pipeline.formatters.gost_formatter import (
    STYLE_MAP, PAGE_SETTINGS, FONT_SETTINGS,
    heading_style_name, body_style_name, list_style_name,
)

if TYPE_CHECKING:
    pass


_TEMPLATE_CT = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
_DOCUMENT_CT = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


def _dotx_to_docx(dotx_bytes: bytes) -> bytes:
    """
    python-docx не умеет открывать .dotx напрямую — тип содержимого
    в [Content_Types].xml отличается от обычного .docx.
    Патчим его в памяти: template.main+xml → document.main+xml.
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(dotx_bytes), "r") as zin:
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(_TEMPLATE_CT, _DOCUMENT_CT)
                zout.writestr(item, data)
    return buf.getvalue()


def _apply_page_settings(doc: Document) -> None:
    """Устанавливает поля страницы A4 по ГОСТ 2.105-2019."""
    sec = doc.sections[0]
    sec.page_width    = Cm(21)    # A4 210 мм
    sec.page_height   = Cm(29.7)  # A4 297 мм
    sec.left_margin   = Cm(3.0)   # ГОСТ: 30 мм
    sec.right_margin  = Cm(1.5)   # ГОСТ: 15 мм
    sec.top_margin    = Cm(2.0)   # ГОСТ: 20 мм
    sec.bottom_margin = Cm(2.0)   # ГОСТ: 20 мм


def _add_page_number(doc: Document) -> None:
    """Добавляет нумерацию страниц внизу по центру."""
    from docx.oxml.ns import qn as _qn
    from docx.oxml import OxmlElement as _OE
    sec = doc.sections[0]
    footer = sec.footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    fld = _OE("w:fldChar")
    fld.set(_qn("w:fldCharType"), "begin")
    run._r.append(fld)
    run2 = para.add_run()
    instrText = _OE("w:instrText")
    instrText.text = "PAGE"
    run2._r.append(instrText)
    run3 = para.add_run()
    fld2 = _OE("w:fldChar")
    fld2.set(_qn("w:fldCharType"), "end")
    run3._r.append(fld2)


def _try_apply_style(doc: Document, paragraph, style_name: str) -> bool:
    """Применяет именованный стиль. Возвращает True при успехе."""
    try:
        paragraph.style = doc.styles[style_name]
        return True
    except KeyError:
        return False


def _apply_heading_fallback(paragraph, level: int) -> None:
    """Программный fallback для заголовков, если стиля .dotx нет."""
    from docx.shared import Pt as _Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _WDA

    para_fmt = paragraph.paragraph_format
    if level == 1:
        paragraph.alignment = _WDA.CENTER
        for run in paragraph.runs:
            run.bold = True
            run.font.size = _Pt(14)
            if paragraph.text:
                run.text = run.text.upper()
    elif level == 2:
        paragraph.alignment = _WDA.LEFT
        para_fmt.first_line_indent = FONT_SETTINGS["first_line"]
        for run in paragraph.runs:
            run.bold = True
            run.font.size = _Pt(14)
    else:
        for run in paragraph.runs:
            run.bold = True
            run.font.size = _Pt(14)


def _apply_body_fallback(paragraph) -> None:
    """Программный fallback для основного текста."""
    from docx.shared import Pt as _Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _WDA

    paragraph.alignment = _WDA.JUSTIFY
    para_fmt = paragraph.paragraph_format
    para_fmt.first_line_indent = FONT_SETTINGS["first_line"]
    para_fmt.line_spacing = FONT_SETTINGS["line_spacing"]
    for run in paragraph.runs:
        run.font.name = FONT_SETTINGS["name"]
        run.font.size = _Pt(FONT_SETTINGS["size_pt"])


class DocxBuilder:
    """
    Собирает .docx из списка IR-элементов.

    Приоритет: открыть .dotx шаблон и применять именованные стили (STYLE_MAP).
    Fallback: программные стили при отсутствии .dotx.
    """

    def __init__(self, dotx_bytes: bytes | None = None) -> None:
        if dotx_bytes:
            self._doc = Document(io.BytesIO(_dotx_to_docx(dotx_bytes)))
        else:
            self._doc = Document()
            _apply_page_settings(self._doc)
            _add_page_number(self._doc)

        self._use_dotx = dotx_bytes is not None

    def build(self, elements: list[IRElement]) -> bytes:
        """Переводит IR-элементы в .docx и возвращает байты."""
        for el in elements:
            match el:
                case IRHeading(text=text, level=level):
                    self._add_heading(text, level)
                case IRParagraph(text=text, style=style):
                    self._add_paragraph(text, style)
                case IRList(items=items, ordered=ordered, style=style):
                    self._add_list(items, ordered, style)
                case IRTable(headers=headers, rows=rows, caption=caption):
                    self._add_table(headers, rows, caption)

        buf = io.BytesIO()
        self._doc.save(buf)
        return buf.getvalue()

    def _add_heading(self, text: str, level: int) -> None:
        p = self._doc.add_paragraph()
        run = p.add_run(text)
        style_name = heading_style_name(level)
        if not _try_apply_style(self._doc, p, style_name):
            _apply_heading_fallback(p, level)

    def _add_paragraph(self, text: str, style: str | None = None) -> None:
        p = self._doc.add_paragraph()
        p.add_run(text)
        resolved = STYLE_MAP.get(style, body_style_name()) if style else body_style_name()
        if not _try_apply_style(self._doc, p, resolved):
            _apply_body_fallback(p)

    def _add_list(self, items: list[str], ordered: bool, style: str | None = None) -> None:
        style_name = STYLE_MAP.get(style) if style else list_style_name(ordered)
        if not style_name:
            style_name = list_style_name(ordered)
        bullet_char = "\u2013 "  # дефис «–» для маркированных
        for i, item in enumerate(items, start=1):
            p = self._doc.add_paragraph()
            if not _try_apply_style(self._doc, p, style_name):
                # Fallback: добавляем символ списка вручную
                prefix = f"{i}. " if ordered else bullet_char
                p.add_run(f"{prefix}{item}")
                _apply_body_fallback(p)
            else:
                p.add_run(item)

    def _add_table(
        self,
        headers: list[str],
        rows: list[list[str]],
        caption: str | None,
    ) -> None:
        if caption:
            cap_p = self._doc.add_paragraph()
            cap_p.add_run(caption)
            cap_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        table = self._doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"

        hdr_row = table.rows[0]
        for i, hdr in enumerate(headers):
            cell = hdr_row.cells[i]
            cell.text = hdr
            _try_apply_style(self._doc, cell.paragraphs[0], STYLE_MAP["table_head"])

        for row_data in rows:
            row = table.add_row()
            for i, val in enumerate(row_data):
                if i < len(row.cells):
                    row.cells[i].text = val
                    _try_apply_style(self._doc, row.cells[i].paragraphs[0], STYLE_MAP["table_cell"])
