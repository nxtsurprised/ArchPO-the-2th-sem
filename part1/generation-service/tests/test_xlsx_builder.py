"""Unit-тесты XlsxBuilder (НМЦК)."""
from __future__ import annotations
import io
import pytest
import openpyxl

from app.pipeline.formatters.xlsx_builder import XlsxBuilder


RATES = {
    "default_rate_per_hour": 3500.0,
    "complexity_coefficients": {
        "low": 1.0,
        "medium": 1.2,
        "high": 1.5,
        "critical": 2.0,
    },
    "overhead_coefficient": 1.15,
    "vat_rate": 0.20,
    "profit_margin": 0.15,
}

SUBSYSTEMS = [
    {"id": "sub-1", "code": "ПБ", "name": "Подсистема ИБ", "order": 1},
    {"id": "sub-2", "code": "АДМ", "name": "Подсистема АДМ", "order": 2},
]

FUNCTIONS = [
    {
        "id": "f1",
        "subsystem_id": "sub-1",
        "code": "ФБ-01",
        "name": "Аутентификация",
        "complexity": "medium",
        "cost_params": {"labor_hours": 40, "rate_per_hour": 3500},
    },
    {
        "id": "f2",
        "subsystem_id": "sub-1",
        "code": "ФБ-02",
        "name": "Журналирование",
        "complexity": "low",
        "cost_params": {"labor_hours": 16, "rate_per_hour": 3500},
    },
    {
        "id": "f3",
        "subsystem_id": "sub-2",
        "code": "АД-01",
        "name": "Управление пользователями",
        "complexity": "high",
        "cost_params": {"labor_hours": 80, "rate_per_hour": 3500},
    },
]


@pytest.fixture
def builder() -> XlsxBuilder:
    return XlsxBuilder()


@pytest.fixture
def workbook(builder):
    raw = builder.build(FUNCTIONS, SUBSYSTEMS, RATES)
    return openpyxl.load_workbook(io.BytesIO(raw))


def test_build_returns_bytes(builder):
    result = builder.build(FUNCTIONS, SUBSYSTEMS, RATES)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_has_three_sheets(workbook):
    assert len(workbook.sheetnames) == 3


def test_nmck_sheet_first(workbook):
    assert workbook.sheetnames[0] == "Расчёт НМЦК"


def test_rates_sheet_exists(workbook):
    assert "Справочник" in workbook.sheetnames


def test_summary_sheet_exists(workbook):
    assert "Сводка" in workbook.sheetnames


def test_nmck_sheet_has_header_row(workbook):
    ws = workbook["Расчёт НМЦК"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, 8)]
    assert headers[0] == "№"
    assert "Функция" in headers
    assert "Трудозатраты, ч" in headers
    assert "Итого, ₽" in headers


def test_nmck_sheet_has_correct_function_count(workbook):
    ws = workbook["Расчёт НМЦК"]
    # Строка 1 — заголовок, строки 2..N — функции
    func_rows = []
    for row_idx in range(2, ws.max_row + 1):
        val = ws.cell(row=row_idx, column=1).value
        if isinstance(val, int):
            func_rows.append(row_idx)
    assert len(func_rows) == len(FUNCTIONS)


def test_nmck_labor_hours_set(workbook):
    ws = workbook["Расчёт НМЦК"]
    # Строка 2: ФБ-01, labor_hours=40
    assert ws.cell(row=2, column=3).value == 40


def test_nmck_itogo_is_formula(workbook):
    ws = workbook["Расчёт НМЦК"]
    # Колонка G строки 2 должна содержать формулу (начинается с =)
    itogo = ws.cell(row=2, column=7).value
    assert isinstance(itogo, str) and itogo.startswith("=")


def test_rates_sheet_base_rate(workbook):
    ws = workbook["Справочник"]
    # B2 — базовая ставка
    assert ws["B2"].value == 3500.0


def test_rates_sheet_vat(workbook):
    ws = workbook["Справочник"]
    # НДС — последняя строка параметров (B9)
    assert ws["B9"].value == 0.20


def test_summary_sheet_has_subsystem_rows(workbook):
    ws = workbook["Сводка"]
    # Строка 1 — заголовок, строки 2+ — подсистемы
    names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert len(names) == len(SUBSYSTEMS)


def test_summary_function_counts(workbook):
    ws = workbook["Сводка"]
    counts = [ws.cell(row=r, column=2).value for r in range(2, 4)]
    # sub-1 — 2 функции, sub-2 — 1 функция
    assert counts[0] == 2
    assert counts[1] == 1


def test_build_empty_functions(builder):
    """Пустой список функций — НМЦК строится без ошибок."""
    result = builder.build([], [], RATES)
    assert isinstance(result, bytes)
    wb = openpyxl.load_workbook(io.BytesIO(result))
    assert "Расчёт НМЦК" in wb.sheetnames


def test_build_from_real_bundle(builder, render_bundle_fixture):
    """Сборка НМЦК из реального render-bundle."""
    functions = render_bundle_fixture["functions"]
    subsystems = render_bundle_fixture["subsystems"]
    rates = render_bundle_fixture["rates"]
    result = builder.build(functions, subsystems, rates)
    assert isinstance(result, bytes)
    wb = openpyxl.load_workbook(io.BytesIO(result))
    ws = wb["Расчёт НМЦК"]
    # Должно быть столько строк с данными, сколько функций
    assert ws.max_row >= len(functions) + 1
