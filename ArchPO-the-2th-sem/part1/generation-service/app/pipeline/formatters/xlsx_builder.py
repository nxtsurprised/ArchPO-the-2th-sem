from __future__ import annotations
import io

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
_BOLD = Font(bold=True, name="Times New Roman", size=12)
_NORMAL = Font(name="Times New Roman", size=12)
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header_style(cell) -> None:
    cell.font = _BOLD
    cell.fill = _HEADER_FILL
    cell.border = _BORDER
    cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")


def _cell_style(cell) -> None:
    cell.font = _NORMAL
    cell.border = _BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="top")


class XlsxBuilder:
    """
    Строит НМЦК в формате .xlsx (openpyxl).

    Лист 1 — Расчёт НМЦК (функции с формулами)
    Лист 2 — Справочник ставок
    Лист 3 — Сводка по подсистемам
    """

    def build(self, functions: list[dict], subsystems: list[dict], rates: dict) -> bytes:
        wb = Workbook()

        self._build_rates_sheet(wb, rates)          # Лист «Справочник»
        self._build_nmck_sheet(wb, functions, rates) # Лист «Расчёт НМЦК»
        self._build_summary_sheet(wb, functions, subsystems)  # Лист «Сводка»

        # Упорядочиваем листы: НМЦК первый
        wb.move_sheet("Расчёт НМЦК", offset=-wb.index(wb["Расчёт НМЦК"]))

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ──────────────────────────────────────────────────────────────────────────
    # Лист 2: Справочник ставок
    # ──────────────────────────────────────────────────────────────────────────
    def _build_rates_sheet(self, wb: Workbook, rates: dict) -> None:
        ws = wb.active
        ws.title = "Справочник"

        params = [
            ("Базовая ставка, ₽/ч",         rates.get("default_rate_per_hour", 0)),
            ("Коэф. сложности Low",          rates.get("complexity_coefficients", {}).get("low", 1.0)),
            ("Коэф. сложности Medium",       rates.get("complexity_coefficients", {}).get("medium", 1.2)),
            ("Коэф. сложности High",         rates.get("complexity_coefficients", {}).get("high", 1.5)),
            ("Коэф. сложности Critical",     rates.get("complexity_coefficients", {}).get("critical", 2.0)),
            ("Накладные расходы (коэф.)",    rates.get("overhead_coefficient", 1.15)),
            ("Маржа",                        rates.get("profit_margin", 0.15)),
            ("НДС",                          rates.get("vat_rate", 0.20)),
        ]

        ws["A1"] = "Параметр"
        ws["B1"] = "Значение"
        _header_style(ws["A1"])
        _header_style(ws["B1"])
        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 15

        for row_idx, (name, value) in enumerate(params, start=2):
            ws[f"A{row_idx}"] = name
            ws[f"B{row_idx}"] = value
            _cell_style(ws[f"A{row_idx}"])
            _cell_style(ws[f"B{row_idx}"])

    # ──────────────────────────────────────────────────────────────────────────
    # Лист 1: Расчёт НМЦК
    # ──────────────────────────────────────────────────────────────────────────
    COMPLEXITY_ROW = {
        "low":      3,   # Справочник!B3
        "medium":   4,   # Справочник!B4
        "high":     5,   # Справочник!B5
        "critical": 6,   # Справочник!B6
    }

    def _build_nmck_sheet(self, wb: Workbook, functions: list[dict], rates: dict) -> None:
        ws = wb.create_sheet("Расчёт НМЦК")

        headers = ["№", "Функция", "Трудозатраты, ч", "Ставка, ₽/ч",
                   "Коэф. сложности", "Накладные", "Итого, ₽"]
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            _header_style(cell)

        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 14
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 12
        ws.column_dimensions["G"].width = 14

        for func_idx, func in enumerate(functions, start=1):
            row = func_idx + 1
            complexity = func.get("complexity", "medium").lower()
            complexity_row = self.COMPLEXITY_ROW.get(complexity, 4)
            labor_hours = func.get("cost_params", {}).get("labor_hours", 0)

            ws.cell(row=row, column=1, value=func_idx)
            ws.cell(row=row, column=2, value=f"{func.get('code', '')} {func.get('name', '')}")
            ws.cell(row=row, column=3, value=labor_hours)
            # Ставка из справочника
            ws.cell(row=row, column=4, value=f"=Справочник!$B$2")
            # Коэф. сложности из справочника
            ws.cell(row=row, column=5, value=f"=Справочник!$B${complexity_row}")
            # Накладные из справочника
            ws.cell(row=row, column=6, value="=Справочник!$B$7")
            # Итого
            c, d, e, f_ = f"C{row}", f"D{row}", f"E{row}", f"F{row}"
            ws.cell(row=row, column=7, value=f"={c}*{d}*{e}*{f_}")

            for col in range(1, 8):
                _cell_style(ws.cell(row=row, column=col))

        # Итоговые строки
        last_data_row = len(functions) + 1
        total_row = last_data_row + 1
        nds_row = total_row + 1
        nmck_row = nds_row + 1

        # Итого без НДС
        ws.cell(row=total_row, column=2, value="Итого без НДС")
        ws.cell(row=total_row, column=3, value=f"=SUM(C2:C{last_data_row})")
        ws.cell(row=total_row, column=7, value=f"=SUM(G2:G{last_data_row})")
        for col in range(1, 8):
            _cell_style(ws.cell(row=total_row, column=col))
        ws.cell(row=total_row, column=2).font = _BOLD
        ws.cell(row=total_row, column=7).font = _BOLD

        # НДС
        ws.cell(row=nds_row, column=2, value="НДС (20%)")
        ws.cell(row=nds_row, column=7, value=f"=G{total_row}*Справочник!$B$9")
        for col in range(1, 8):
            _cell_style(ws.cell(row=nds_row, column=col))

        # НМЦК
        ws.cell(row=nmck_row, column=2, value="НМЦК")
        ws.cell(row=nmck_row, column=7, value=f"=G{total_row}+G{nds_row}")
        for col in range(1, 8):
            cell = ws.cell(row=nmck_row, column=col)
            _cell_style(cell)
            cell.font = _BOLD

    # ──────────────────────────────────────────────────────────────────────────
    # Лист 3: Сводка по подсистемам
    # ──────────────────────────────────────────────────────────────────────────
    def _build_summary_sheet(self, wb: Workbook, functions: list[dict], subsystems: list[dict]) -> None:
        ws = wb.create_sheet("Сводка")

        headers = ["Подсистема", "Кол-во функций", "Стоимость, ₽"]
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            _header_style(cell)

        ws.column_dimensions["A"].width = 40
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 18

        # Кол-во функций и стоимость считаем из данных (не ссылками на другой лист,
        # т.к. строки подсистем не фиксированы) — используем статические формулы
        sub_map: dict[str, dict] = {s["id"]: s for s in subsystems}
        sub_counts: dict[str, int] = {}
        sub_hours: dict[str, float] = {}

        for func in functions:
            sid = func.get("subsystem_id", "")
            sub_counts[sid] = sub_counts.get(sid, 0) + 1
            sub_hours[sid] = sub_hours.get(sid, 0.0) + func.get("cost_params", {}).get("labor_hours", 0)

        for row_idx, sub in enumerate(sorted(subsystems, key=lambda s: s.get("order", 0)), start=2):
            sid = sub["id"]
            ws.cell(row=row_idx, column=1, value=f"{sub.get('code', '')} — {sub.get('name', '')}")
            ws.cell(row=row_idx, column=2, value=sub_counts.get(sid, 0))
            # Стоимость = трудозатраты × базовая_ставка (упрощённая сводка)
            ws.cell(
                row=row_idx, column=3,
                value=f"={sub_hours.get(sid, 0)}*Справочник!$B$2",
            )
            for col in range(1, 4):
                _cell_style(ws.cell(row=row_idx, column=col))
