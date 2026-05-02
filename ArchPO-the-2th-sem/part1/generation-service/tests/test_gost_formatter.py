"""Unit-тесты DocxBuilder (ГОСТ 2.105 форматирование) и GostFormatter."""
from __future__ import annotations
import io
import pytest

from app.pipeline.formatters.gost_formatter import (
    STYLE_MAP, PAGE_SETTINGS, FONT_SETTINGS, heading_style_name, body_style_name, list_style_name,
)
from app.pipeline.formatters.docx_builder import DocxBuilder
from app.pipeline.renderers.manual_renderer import IRHeading, IRParagraph, IRList, IRTable


# ── GostFormatter / STYLE_MAP ─────────────────────────────────────────────────

def test_style_map_has_required_keys():
    required = {"heading_1", "heading_2", "heading_3", "body_text", "list_bullet", "list_number", "table_head", "table_cell"}
    assert required.issubset(STYLE_MAP.keys())


def test_heading_style_names():
    assert heading_style_name(1) == "Заголовок 1"
    assert heading_style_name(2) == "Заголовок 2"
    assert heading_style_name(3) == "Заголовок 3"


def test_list_style_name_ordered():
    assert list_style_name(True) == STYLE_MAP["list_number"]


def test_list_style_name_bullet():
    assert list_style_name(False) == STYLE_MAP["list_bullet"]


def test_page_settings_a4():
    assert PAGE_SETTINGS["width"] == 11906
    assert PAGE_SETTINGS["height"] == 16838
    assert PAGE_SETTINGS["left"] == 1701   # 30 мм
    assert PAGE_SETTINGS["right"] == 850   # 15 мм


def test_font_settings_tnr():
    assert FONT_SETTINGS["name"] == "Times New Roman"
    assert FONT_SETTINGS["size_pt"] == 14


# ── DocxBuilder — fallback без .dotx ─────────────────────────────────────────

@pytest.fixture
def builder_no_dotx() -> DocxBuilder:
    return DocxBuilder(dotx_bytes=None)


def test_build_returns_bytes(builder_no_dotx):
    elements = [IRHeading(text="Раздел 1", level=1)]
    result = builder_no_dotx.build(elements)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_build_heading_in_doc(builder_no_dotx):
    from docx import Document
    elements = [IRHeading(text="ТЕСТОВЫЙ ЗАГОЛОВОК", level=1)]
    doc_bytes = builder_no_dotx.build(elements)
    doc = Document(io.BytesIO(doc_bytes))
    texts = [p.text for p in doc.paragraphs]
    assert any("ТЕСТОВЫЙ ЗАГОЛОВОК" in t for t in texts)


def test_build_paragraph_in_doc(builder_no_dotx):
    from docx import Document
    elements = [IRParagraph(text="Тестовый абзац с содержимым.")]
    doc_bytes = builder_no_dotx.build(elements)
    doc = Document(io.BytesIO(doc_bytes))
    texts = " ".join(p.text for p in doc.paragraphs)
    assert "Тестовый абзац" in texts


def test_build_ordered_list(builder_no_dotx):
    from docx import Document
    items = ["Первый элемент", "Второй элемент", "Третий элемент"]
    elements = [IRList(items=items, ordered=True)]
    doc_bytes = builder_no_dotx.build(elements)
    doc = Document(io.BytesIO(doc_bytes))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Первый элемент" in all_text
    assert "Второй элемент" in all_text


def test_build_bullet_list(builder_no_dotx):
    from docx import Document
    items = ["Пункт А", "Пункт Б"]
    elements = [IRList(items=items, ordered=False)]
    doc_bytes = builder_no_dotx.build(elements)
    doc = Document(io.BytesIO(doc_bytes))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Пункт А" in all_text


def test_build_table(builder_no_dotx):
    from docx import Document
    elements = [
        IRTable(
            headers=["Колонка 1", "Колонка 2"],
            rows=[["Строка 1, ячейка 1", "Строка 1, ячейка 2"]],
            caption="Таблица 1 — Тест",
        )
    ]
    doc_bytes = builder_no_dotx.build(elements)
    doc = Document(io.BytesIO(doc_bytes))
    assert len(doc.tables) == 1
    tbl = doc.tables[0]
    assert tbl.rows[0].cells[0].text == "Колонка 1"
    assert tbl.rows[1].cells[1].text == "Строка 1, ячейка 2"


def test_build_mixed_elements(builder_no_dotx):
    """Полный документ с заголовком, абзацем, списком и таблицей."""
    from docx import Document
    elements = [
        IRHeading(text="1 Общие сведения", level=1),
        IRParagraph(text="Заказчик: Минстрой России"),
        IRList(items=["Первое", "Второе"], ordered=False),
        IRTable(headers=["№", "Название"], rows=[["1", "Авторизация"]], caption=None),
    ]
    doc_bytes = builder_no_dotx.build(elements)
    doc = Document(io.BytesIO(doc_bytes))
    assert len(doc.paragraphs) >= 3
    assert len(doc.tables) == 1


def test_build_from_real_bundle(render_bundle_fixture):
    """Сборка .docx из реального render-bundle без ошибок."""
    from app.pipeline.section_renderer import render_section

    template = render_bundle_fixture.get("template", {})
    all_elements = []
    for section in template.get("sections", []):
        elements = render_section(section, render_bundle_fixture)
        all_elements.extend(elements)

    builder = DocxBuilder(dotx_bytes=None)
    doc_bytes = builder.build(all_elements)
    assert isinstance(doc_bytes, bytes)
    assert len(doc_bytes) > 500  # валидный docx весит хотя бы несколько сотен байт
